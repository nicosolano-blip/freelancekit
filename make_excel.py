from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "50 Ideas"

# ── Colors ────────────────────────────────────────────────────────────────────
NAVY   = "1B3A6B"
TEAL   = "0E7490"
WHITE  = "FFFFFF"
GRAY1  = "F1F5F9"
GRAY2  = "E2E8F0"
DARK   = "1E293B"
MID    = "475569"
GREEN  = "16A34A"
ORANGE = "EA580C"

CAT_COLORS = {
    "Freelancers y Autónomos":       ("EFF6FF", "1D4ED8"),
    "Pequeños Negocios":             ("F0FDF4", "15803D"),
    "Construcción y Oficios":        ("FFF7ED", "C2410C"),
    "Inmobiliario y Finanzas":       ("FDF4FF", "7E22CE"),
    "Salud y Bienestar":             ("F0FDFA", "0F766E"),
    "Educación y Productividad":     ("FEFCE8", "A16207"),
    "Autos y Transporte":            ("FEF2F2", "B91C1C"),
    "Recursos Humanos y Empleo":     ("F5F3FF", "6D28D9"),
    "Eventos y Vida Cotidiana":      ("FFF1F2", "BE123C"),
}

ideas = [
    # Freelancers y Autónomos
    ("Freelancers y Autónomos", "Calculadora de Tarifa Hora Freelancer",
     "Calcula la tarifa óptima por hora según gastos mensuales, impuestos, días trabajados e ingreso deseado.",
     "3–5 hs"),
    ("Freelancers y Autónomos", "Generador de Cotización/Presupuesto PDF",
     "Formulario donde el usuario ingresa datos del proyecto y genera un PDF profesional listo para enviar al cliente.",
     "6–10 hs"),
    ("Freelancers y Autónomos", "Calculadora de Tarifa en USD para LATAM",
     "Estima cuánto cobrar en dólares considerando paridad de poder adquisitivo, tipo de cliente y posicionamiento.",
     "3–4 hs"),
    ("Freelancers y Autónomos", "Generador de Contrato Simple",
     "Crea contratos básicos personalizables para servicios freelance. Exportable en PDF o texto.",
     "5–8 hs"),
    ("Freelancers y Autónomos", "Tracker de Horas y Facturación Mensual",
     "Registra horas por proyecto en el browser, calcula total a facturar y genera resumen mensual.",
     "6–10 hs"),
    ("Freelancers y Autónomos", "Calculadora de Aumento de Tarifa por Inflación",
     "Indica cuánto subir los precios para mantener el mismo ingreso real dado un porcentaje de inflación.",
     "2–3 hs"),

    # Pequeños Negocios
    ("Pequeños Negocios", "Calculadora de Punto de Equilibrio",
     "Calcula cuántas unidades o cuánto revenue necesita un negocio para cubrir costos fijos y variables.",
     "3–5 hs"),
    ("Pequeños Negocios", "Calculadora de Margen Real para E-commerce",
     "Precio de venta menos costo, comisión de plataforma, shipping e impuestos. Muestra ganancia neta real.",
     "4–6 hs"),
    ("Pequeños Negocios", "Estimador de ROAS para Campañas Digitales",
     "Calcula qué ROAS mínimo necesitás en Meta/Google para que la publicidad sea rentable según tu margen.",
     "3–4 hs"),
    ("Pequeños Negocios", "Calculadora de CAC (Costo de Adquisición de Cliente)",
     "Divide inversión en marketing entre clientes obtenidos y compara contra el LTV estimado del cliente.",
     "3–4 hs"),
    ("Pequeños Negocios", "Generador de Flujo de Caja Básico",
     "Plantilla interactiva mensual de ingresos y egresos. Exporta a Excel. Sin macros, sin cuentas.",
     "5–8 hs"),
    ("Pequeños Negocios", "Simulador de Escenarios de Negocio",
     "Qué pasa con las ganancias si subo precios X%, si vendo Y% más o si bajo costos Z%. Tres escenarios en paralelo.",
     "5–7 hs"),
    ("Pequeños Negocios", "Calculadora de Precio de Producto Físico/Artesanal",
     "Suma materiales + tiempo + overhead + ganancia deseada para obtener el precio de venta correcto.",
     "3–5 hs"),
    ("Pequeños Negocios", "Generador de Política de Privacidad para LATAM",
     "Genera términos y condiciones y política de privacidad adaptados a legislación latinoamericana.",
     "4–6 hs"),

    # Construcción y Oficios
    ("Construcción y Oficios", "Calculadora de Materiales por m²",
     "Calcula ladrillos, cemento, arena y agua necesarios para una superficie dada. Incluye % de desperdicio.",
     "3–5 hs"),
    ("Construcción y Oficios", "Estimador de Presupuesto de Obra/Refacción",
     "Habitación por habitación: seleccionás tipo de trabajo, m² y muestra costo estimado de materiales + mano de obra.",
     "8–12 hs"),
    ("Construcción y Oficios", "Calculadora de Pintura",
     "Litros necesarios según m², tipo de superficie (lisa/rugosa), número de manos y rendimiento del producto.",
     "2–3 hs"),
    ("Construcción y Oficios", "Calculadora de Pisos y Cerámicos",
     "M² necesarios más porcentaje de desperdicio por tipo de colocación (diagonal, recto, combinado).",
     "2–3 hs"),
    ("Construcción y Oficios", "Generador de Presupuesto de Obra Descargable",
     "El contratista ingresa ítems, cantidades y precios. Genera PDF de presupuesto con logo y datos de empresa.",
     "8–14 hs"),
    ("Construcción y Oficios", "Calculadora de Sección de Cable Eléctrico",
     "Determina la sección correcta de cable según consumo en watts, distancia y tensión. Con tabla de referencia.",
     "4–6 hs"),

    # Inmobiliario y Finanzas
    ("Inmobiliario y Finanzas", "Calculadora de ROI de Alquiler",
     "Ingresos por alquiler menos expensas, impuestos, mantenimiento y vacancia estimada. ROI anual neto.",
     "4–6 hs"),
    ("Inmobiliario y Finanzas", "Estimador de Gastos al Comprar una Propiedad",
     "Suma escribanía, impuesto de sellos, comisión inmobiliaria y otros gastos ocultos al precio de compra.",
     "3–4 hs"),
    ("Inmobiliario y Finanzas", "Calculadora de Costo Real de un Crédito Hipotecario",
     "Capital + intereses + seguros + gastos admin a lo largo de los años. Muestra el costo total real pagado.",
     "4–6 hs"),
    ("Inmobiliario y Finanzas", "Comparador Alquiler vs Compra",
     "En qué mes/año conviene más comprar que alquilar dado precio del inmueble, alquiler, inflación y plazo.",
     "5–8 hs"),
    ("Inmobiliario y Finanzas", "Calculadora de Inflación Real",
     "Cuánto vale hoy lo que pagué hace X meses/años según inflación oficial o estimada. Especialmente para ARG.",
     "2–3 hs"),
    ("Inmobiliario y Finanzas", "Simulador de Ahorro: Pesos vs Dólares vs Plazo Fijo",
     "Compara el valor real de ahorrar en distintos instrumentos dado el plazo y la inflación esperada.",
     "4–6 hs"),

    # Salud y Bienestar
    ("Salud y Bienestar", "Calculadora de Macros con Alimentos Locales",
     "Calcula proteínas, carbohidratos y grasas diarias según objetivo. Base de datos con alimentos de LATAM.",
     "8–14 hs"),
    ("Salud y Bienestar", "Planificador de Rutina de Ejercicio Descargable",
     "El usuario elige nivel, objetivo y días disponibles. Genera rutina semanal en PDF para imprimir.",
     "5–8 hs"),
    ("Salud y Bienestar", "Calculadora de Sueño por Ciclos",
     "A qué hora dormirse para despertar descansado según ciclos de 90 minutos y hora de levantarse deseada.",
     "2–3 hs"),
    ("Salud y Bienestar", "Calculadora de Hidratación Diaria",
     "Cantidad de agua recomendada según peso, nivel de actividad física y temperatura ambiente.",
     "2–3 hs"),
    ("Salud y Bienestar", "Estimador de Calorías en Actividades Cotidianas",
     "Cuántas calorías quemás caminando, subiendo escaleras, limpiando, etc. Sin registro ni login.",
     "3–4 hs"),
    ("Salud y Bienestar", "Calculadora de Costo de Hábitos Anuales",
     "Cuánto gastás por año en café diario, cigarrillos, delivery, etc. Muy emocional y muy viral.",
     "3–4 hs"),

    # Educación y Productividad
    ("Educación y Productividad", "Temporizador Pomodoro con Estadísticas",
     "Temporizador 25/5 min con registro de sesiones del día, racha de días y resumen semanal. Sin login.",
     "4–6 hs"),
    ("Educación y Productividad", "Conversor de Calificaciones entre Sistemas Educativos",
     "Convierte notas entre sistema argentino, americano, europeo y otros. Útil para quienes aplican al exterior.",
     "3–4 hs"),
    ("Educación y Productividad", "Calculadora de Promedio Universitario Ponderado",
     "El alumno ingresa materias, notas y créditos/horas. Calcula promedio ponderado y cuánto necesita en el final.",
     "3–5 hs"),
    ("Educación y Productividad", "Generador de Plan de Estudio",
     "Ingresás fecha de examen y temas a estudiar. Genera calendario de estudio distribuido en días disponibles.",
     "5–8 hs"),
    ("Educación y Productividad", "Estimador de Tiempo de Lectura",
     "Pegás texto o ingresás páginas y palabras por minuto. Calcula cuánto tardás en leer el material.",
     "1–2 hs"),

    # Autos y Transporte
    ("Autos y Transporte", "Calculadora de Costo Real de Tener un Auto",
     "Suma patente, seguro, nafta, service, devaluación y estacionamiento. Muestra el costo mensual y anual real.",
     "4–6 hs"),
    ("Autos y Transporte", "Comparador Auto Propio vs Uber/Taxi/Remis",
     "Según kilómetros mensuales, calcula en qué punto conviene más tener auto propio versus usar servicios.",
     "4–5 hs"),
    ("Autos y Transporte", "Estimador de Consumo de Combustible por Viaje",
     "Distancia × consumo del vehículo × precio del combustible = costo total del viaje en auto.",
     "2–3 hs"),
    ("Autos y Transporte", "Calculadora de Cuándo Vender el Auto",
     "Según antigüedad, km, valor actual y costo de reparaciones proyectadas, indica si conviene vender o mantener.",
     "4–6 hs"),

    # Recursos Humanos y Empleo
    ("Recursos Humanos y Empleo", "Calculadora de Costo de Empleado para el Empleador",
     "Salario bruto + cargas sociales + obra social + ART + SAC. Muestra el costo laboral total real mensual.",
     "4–6 hs"),
    ("Recursos Humanos y Empleo", "Calculadora de Liquidación de Sueldo",
     "Para empleados que quieren verificar si su recibo está bien. Ingresás categoría y sueldo, te muestra los descuentos.",
     "5–8 hs"),
    ("Recursos Humanos y Empleo", "Estimador de Indemnización por Despido",
     "Según antigüedad, sueldo y tipo de despido calcula la indemnización que corresponde legalmente.",
     "3–5 hs"),
    ("Recursos Humanos y Empleo", "Comparador de Ofertas de Trabajo",
     "Dos ofertas en paralelo: sueldo neto, beneficios, distancia, horario. Puntúa cuál conviene más en términos reales.",
     "5–7 hs"),

    # Eventos y Vida Cotidiana
    ("Eventos y Vida Cotidiana", "Calculadora de Costo por Persona para Eventos",
     "Catering, salón, música, flores, decoración dividido entre invitados. Muestra el costo real por persona.",
     "3–4 hs"),
    ("Eventos y Vida Cotidiana", "Divisor de Gastos en Viaje Grupal",
     "Quién puso qué, quién debe a quién. Más simple que Splitwise, sin app ni registro. 100% en el browser.",
     "5–8 hs"),
    ("Eventos y Vida Cotidiana", "Calculadora de Conveniencia de Mudanza",
     "Contratar flete vs alquilar camioneta vs pagar un servicio completo. Compara costos según distancia y volumen.",
     "3–4 hs"),
    ("Eventos y Vida Cotidiana", "Estimador de Costo Anual de Mascota",
     "Por especie y tamaño: comida, veterinario, vacunas, higiene, accesorios. Para que la gente calcule antes de adoptar.",
     "3–5 hs"),
    ("Eventos y Vida Cotidiana", "Comparador de Costo Semanal de Dietas",
     "Keto vs vegana vs mediterránea vs low-cost. Costo real en alimentos por semana según dieta elegida.",
     "5–7 hs"),
]

# ── Sheet setup ───────────────────────────────────────────────────────────────
thin = Side(style="thin", color="D1D5DB")
med  = Side(style="medium", color="CBD5E1")

def cell_border(top=thin, bottom=thin, left=thin, right=thin):
    return Border(top=top, bottom=bottom, left=left, right=right)

# Header row
headers = ["#", "Categoría", "Nombre del Proyecto", "Descripción", "Tiempo Estimado (con Claude)"]
col_widths = [4, 26, 36, 72, 28]

for col, (h, w) in enumerate(zip(headers, col_widths), 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color=WHITE, name="Arial", size=11)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = cell_border()
    ws.column_dimensions[get_column_letter(col)].width = w

ws.row_dimensions[1].height = 28

# Data rows
for i, (cat, name, desc, time) in enumerate(ideas, 1):
    row = i + 1
    bg, accent = CAT_COLORS[cat]
    alt = "F8FAFC" if i % 2 == 0 else "FFFFFF"

    vals = [i, cat, name, desc, time]
    for col, val in enumerate(vals, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = Font(name="Arial", size=10,
                      color=accent if col == 2 else (DARK if col in (1, 3, 5) else MID),
                      bold=(col in (1, 3, 5)))
        c.fill = PatternFill("solid", fgColor=bg if col == 2 else alt)
        c.alignment = Alignment(vertical="top", wrap_text=True,
                                horizontal="center" if col in (1, 5) else "left")
        c.border = cell_border()
    ws.row_dimensions[row].height = 52

# Freeze header
ws.freeze_panes = "A2"

# Auto-filter
ws.auto_filter.ref = f"A1:E{len(ideas)+1}"

# ── Summary sheet ─────────────────────────────────────────────────────────────
ws2 = wb.create_sheet("Resumen por Categoría")

ws2.column_dimensions["A"].width = 36
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 32

for col, h in enumerate(["Categoría", "# Proyectos", "Tiempo Total Estimado"], 1):
    c = ws2.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color=WHITE, name="Arial", size=11)
    c.fill = PatternFill("solid", fgColor=TEAL)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = cell_border()

from collections import Counter
cat_counts = Counter(cat for cat, *_ in ideas)
cat_times = {}
for cat, name, desc, time in ideas:
    if cat not in cat_times:
        cat_times[cat] = []
    # parse min/max
    parts = time.replace(" hs", "").split("–")
    cat_times[cat].append((int(parts[0]), int(parts[1])))

for r, (cat, count) in enumerate(cat_counts.items(), 2):
    times = cat_times[cat]
    total_min = sum(t[0] for t in times)
    total_max = sum(t[1] for t in times)
    bg, accent = CAT_COLORS[cat]
    alt = "F8FAFC" if r % 2 == 0 else "FFFFFF"

    for col, val in enumerate([cat, count, f"{total_min}–{total_max} hs"], 1):
        c = ws2.cell(row=r, column=col, value=val)
        c.font = Font(name="Arial", size=10, bold=(col==1), color=accent if col==1 else DARK)
        c.fill = PatternFill("solid", fgColor=bg if col==1 else alt)
        c.alignment = Alignment(vertical="center", horizontal="center" if col==2 else "left")
        c.border = cell_border()
    ws2.row_dimensions[r].height = 22

# Total row
total_row = len(cat_counts) + 2
all_times = [(int(t.replace(" hs","").split("–")[0]), int(t.replace(" hs","").split("–")[1])) for _,_,_,t in ideas]
grand_min = sum(t[0] for t in all_times)
grand_max = sum(t[1] for t in all_times)

for col, val in enumerate(["TOTAL", 50, f"{grand_min}–{grand_max} hs"], 1):
    c = ws2.cell(row=total_row, column=col, value=val)
    c.font = Font(bold=True, color=WHITE, name="Arial", size=11)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = cell_border(top=med, bottom=med, left=med, right=med)
ws2.row_dimensions[total_row].height = 24

ws2.freeze_panes = "A2"

wb.save("/sessions/sweet-sharp-ptolemy/mnt/outputs/50_Ideas_Apps_Web.xlsx")
print("Done")
